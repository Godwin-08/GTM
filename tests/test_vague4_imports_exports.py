"""Tests unitaires et d'intégration spécifiques à la Vague 4 : Imports & Exports.
- Point 6 : Export Utilisateurs & Formateurs (CSV UTF-8 BOM et Excel XLSX stylisé)
- Point 8 : Import par lot de participants (CSV/Excel) avec validation granulaire et rapport d'erreurs
"""

import io
import unittest
import openpyxl
from werkzeug.security import generate_password_hash
from app import create_app
from app.config import Config
from app.extensions import db
from app.models import Utilisateur, Role, Formateur, Domaine, Client, Participant


class Vague4ImportsExportsTestCase(unittest.TestCase):
    """Suite de validation pour les imports et exports de la Vague 4."""

    @classmethod
    def setUpClass(cls):
        cls.original_database_uri = Config.SQLALCHEMY_DATABASE_URI
        Config.SQLALCHEMY_DATABASE_URI = "sqlite://"
        cls.app = create_app()
        cls.app.config.update(TESTING=True)

    @classmethod
    def tearDownClass(cls):
        Config.SQLALCHEMY_DATABASE_URI = cls.original_database_uri

    def setUp(self):
        self.context = self.app.app_context()
        self.context.push()
        db.drop_all()
        db.create_all()

        self.admin_role = Role(nom="admin")
        self.gestionnaire_role = Role(nom="gestionnaire")
        self.formateur_role = Role(nom="formateur")
        db.session.add_all([self.admin_role, self.gestionnaire_role, self.formateur_role])
        db.session.flush()

        self.domaine = Domaine(nom="Management Agile")
        db.session.add(self.domaine)
        db.session.flush()

        self.admin = Utilisateur(
            nom="Admin User",
            email="admin.v4@test.ma",
            mot_de_passe_hash=generate_password_hash("AdminPass123!"),
            role=self.admin_role,
            actif=True,
        )
        self.user_formateur = Utilisateur(
            nom="Formateur User",
            email="formateur.v4@test.ma",
            mot_de_passe_hash=generate_password_hash("FormateurPass123!"),
            role=self.formateur_role,
            actif=True,
        )
        self.client_corp = Client(
            nom_entreprise="OCP Group",
            contact_email="rh@ocp.ma",
            secteur="Industrie",
        )
        db.session.add_all([self.admin, self.user_formateur, self.client_corp])
        db.session.flush()

        self.formateur_profil = Formateur(
            nom="Nadia Mansouri",
            email="nadia.mansouri@test.ma",
            domaine=self.domaine,
            utilisateur=self.user_formateur,
        )
        db.session.add(self.formateur_profil)
        db.session.commit()

        self.client = self.app.test_client()

    def tearDown(self):
        db.session.remove()
        self.context.pop()

    def connecter(self, user):
        with self.client.session_transaction() as session:
            session["_user_id"] = str(user.id)
            session["_fresh"] = True

    # -------------------------------------------------------------
    # Point 6 : Tests des Exports Utilisateurs & Formateurs
    # -------------------------------------------------------------
    def test_export_utilisateurs_csv_succes(self):
        """Vérifie la génération du flux CSV des utilisateurs avec BOM UTF-8."""
        self.connecter(self.admin)
        res = self.client.get("/api/utilisateurs/export/csv")
        self.assertEqual(res.status_code, 200)
        self.assertIn("text/csv", res.headers.get("Content-Type", ""))
        self.assertTrue(res.data.startswith(b"\xef\xbb\xbf"))  # BOM UTF-8
        contenu = res.data.decode("utf-8-sig")
        self.assertIn("Nom complet", contenu)
        self.assertIn("admin.v4@test.ma", contenu)
        self.assertIn("Formateur User", contenu)

    def test_export_utilisateurs_xlsx_succes(self):
        """Vérifie la génération d'un classeur Excel valide pour les utilisateurs."""
        self.connecter(self.admin)
        res = self.client.get("/api/utilisateurs/export/xlsx")
        self.assertEqual(res.status_code, 200)
        self.assertIn("spreadsheetml", res.headers.get("Content-Type", ""))
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Utilisateurs", wb.sheetnames)
        ws = wb["Utilisateurs"]
        self.assertEqual(ws.cell(row=1, column=2).value, "Nom complet")

    def test_export_formateurs_csv_succes(self):
        """Vérifie l'export CSV des formateurs avec métriques."""
        self.connecter(self.admin)
        res = self.client.get("/api/formateurs/export/csv")
        self.assertEqual(res.status_code, 200)
        self.assertIn("text/csv", res.headers.get("Content-Type", ""))
        contenu = res.data.decode("utf-8-sig")
        self.assertIn("Nadia Mansouri", contenu)
        self.assertIn("Management Agile", contenu)
        self.assertIn("Interne", contenu)

    def test_export_formateurs_xlsx_succes(self):
        """Vérifie l'export Excel XLSX des formateurs."""
        self.connecter(self.admin)
        res = self.client.get("/api/formateurs/export/xlsx")
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Formateurs", wb.sheetnames)

    def test_exports_interdits_au_formateur(self):
        """Le rôle Formateur ne doit pas pouvoir exporter les utilisateurs ni l'annuaire formateurs."""
        self.connecter(self.user_formateur)
        self.assertEqual(self.client.get("/api/utilisateurs/export/csv").status_code, 403)
        self.assertEqual(self.client.get("/api/utilisateurs/export/xlsx").status_code, 403)
        self.assertEqual(self.client.get("/api/formateurs/export/csv").status_code, 403)
        self.assertEqual(self.client.get("/api/formateurs/export/xlsx").status_code, 403)

    # -------------------------------------------------------------
    # Point 8 : Tests de l'Import de Participants
    # -------------------------------------------------------------
    def test_telechargement_template_import(self):
        """Le template d'import CSV doit être téléchargeable."""
        self.connecter(self.admin)
        res = self.client.get("/api/participants/import/template")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.data.startswith(b"\xef\xbb\xbf"))
        contenu = res.data.decode("utf-8-sig")
        self.assertIn("Nom", contenu)
        self.assertIn("Email", contenu)
        self.assertIn("Entreprise", contenu)

    def test_import_csv_virgule_succes(self):
        """Import d'un fichier CSV avec séparateur virgule."""
        self.connecter(self.admin)
        csv_content = (
            "Nom,Email,Entreprise\n"
            "Said Tazi,said.tazi@ocp.ma,OCP Group\n"
            "Houda Rami,houda.rami@ocp.ma,OCP Group\n"
        ).encode("utf-8")

        res = self.client.post(
            "/api/participants/import",
            data={"fichier": (io.BytesIO(csv_content), "participants.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["succes"])
        self.assertEqual(data["nb_importes"], 2)
        self.assertEqual(data["nb_erreurs"], 0)

        # Vérification en base
        p1 = Participant.query.filter_by(email="said.tazi@ocp.ma").first()
        self.assertIsNotNone(p1)
        self.assertEqual(p1.nom, "Said Tazi")
        self.assertEqual(p1.client_id, self.client_corp.id)

    def test_import_csv_point_virgule_succes(self):
        """Import d'un fichier CSV avec séparateur point-virgule (dialecte Excel FR)."""
        self.connecter(self.admin)
        csv_content = (
            "Nom;Email;Entreprise\n"
            "Anas Berrada;anas.berrada@ocp.ma;OCP Group\n"
        ).encode("utf-8-sig")

        res = self.client.post(
            "/api/participants/import",
            data={"fichier": (io.BytesIO(csv_content), "liste_participants.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.get_json()["nb_importes"], 1)

    def test_import_xlsx_succes(self):
        """Import depuis un classeur Excel .xlsx."""
        self.connecter(self.admin)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Nom complet", "Courriel", "Société"])
        ws.append(["Tariq Hilali", "tariq.hilali@ocp.ma", "OCP Group"])
        xlsx_buffer = io.BytesIO()
        wb.save(xlsx_buffer)
        xlsx_buffer.seek(0)

        res = self.client.post(
            "/api/participants/import",
            data={"fichier": (xlsx_buffer, "participants.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.get_json()["nb_importes"], 1)
        self.assertIsNotNone(Participant.query.filter_by(email="tariq.hilali@ocp.ma").first())

    def test_import_rapport_granulaire_avec_erreurs(self):
        """Import mixte : lignes valides enregistrées, lignes invalides recensées avec motif."""
        self.connecter(self.admin)
        csv_content = (
            "Nom,Email,Entreprise\n"
            "Participant Valide,valide@ocp.ma,OCP Group\n"
            ",sans.nom@ocp.ma,OCP Group\n"
            "Email Invalide,invalide-email-format,OCP Group\n"
            "Client Inconnu,inconnu@test.ma,Entreprise Fantome Non Existante\n"
        ).encode("utf-8")

        res = self.client.post(
            "/api/participants/import",
            data={"fichier": (io.BytesIO(csv_content), "test_mixte.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertEqual(data["total_lignes"], 4)
        self.assertEqual(data["nb_importes"], 1)
        self.assertEqual(data["nb_erreurs"], 3)
        self.assertEqual(len(data["erreurs"]), 3)

        # Ligne valide créée
        self.assertIsNotNone(Participant.query.filter_by(email="valide@ocp.ma").first())

    def test_import_interdit_au_formateur(self):
        """Un formateur ne peut pas importer de participants (403)."""
        self.connecter(self.user_formateur)
        res = self.client.post(
            "/api/participants/import",
            data={"fichier": (io.BytesIO(b"data"), "test.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 403)

