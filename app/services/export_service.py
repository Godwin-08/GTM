"""
==============================================================================
Service de Génération des Exports (CSV, Excel XLSX & Rapports PDF)
==============================================================================
Ce module gère l'ensemble des fonctionnalités d'exportation de Galaxy Solutions :
1. CSV : Format universel encodé en UTF-8 avec BOM (utf-8-sig) pour ouverture
   immédiate et sans altération d'accents dans Microsoft Excel.
2. Excel (.xlsx) : Génération de classeurs stylisés via openpyxl respectant
   la charte graphique Emerald Graphite (en-têtes émeraude, lignes alternées,
   largeurs automatiques de colonnes).
3. PDF ReportLab :
   - Rapport de Synthèse & Pilotage Décisionnel A4 (KPIs et Points d'attention).
   - Fiche de Session & Feuille d'Émargement officielle (avec zones d'émargement matin/après-midi).
"""

import csv
import io
from datetime import datetime
from flask import Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable


# =============================================================================
# 1. EXPORT CSV (UTF-8 avec BOM)
# =============================================================================

def generer_csv_response(nom_fichier, en_tetes, lignes):
    """
    Génère un flux binaire CSV téléchargeable avec en-têtes et encodage UTF-8 + BOM.
    
    :param nom_fichier: Nom du fichier à télécharger (ex: 'formations_export.csv')
    :param en_tetes: Dictionnaire {cle_dict: 'Libellé En-tête'} ou liste de clés
    :param lignes: Liste de dictionnaires contenant les données
    :return: Objet Response Flask configuré avec Content-Disposition attachment
    """
    output_bytes = io.BytesIO()
    # utf-8-sig insère le marqueur Byte Order Mark (BOM) indispensable pour Excel Windows
    text_stream = io.TextIOWrapper(output_bytes, encoding="utf-8-sig", newline="")

    if isinstance(en_tetes, dict):
        fieldnames = list(en_tetes.keys())
        header_map = en_tetes
    else:
        fieldnames = list(en_tetes)
        header_map = {f: f for f in fieldnames}

    writer = csv.DictWriter(text_stream, fieldnames=fieldnames, extrasaction="ignore")
    # Écriture de la ligne d'en-tête personnalisée
    writer.writerow(header_map)

    # Écriture des lignes de données
    for ligne in lignes:
        writer.writerow(ligne)

    text_stream.flush()
    csv_data = output_bytes.getvalue()

    response = Response(csv_data, mimetype="text/csv; charset=utf-8")
    response.headers["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response


# =============================================================================
# 2. EXPORT EXCEL (XLSX Stylisé Emerald Graphite)
# =============================================================================

def generer_xlsx_response(nom_fichier, en_tetes, lignes, titre_feuille="Données"):
    """
    Génère un classeur Excel .xlsx avec mise en forme professionnelle :
    - En-tête vert émeraude (#047857) avec texte blanc en gras.
    - Lignes alternées avec fond ardoise clair (#F8FAFC).
    - Alignement automatique (nombres à droite, dates au centre, textes à gauche).
    - Calcul automatique de la largeur des colonnes.
    
    :param nom_fichier: Nom du fichier XLSX cible (ex: 'sessions.xlsx')
    :param en_tetes: Dictionnaire {cle: 'Libellé'} ou liste
    :param lignes: Liste de dictionnaires de données
    :param titre_feuille: Nom de l'onglet Excel (limité à 31 caractères)
    :return: Objet Response Flask
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titre_feuille[:31]  # Contrainte Excel : nom d'onglet <= 31 car.

    if isinstance(en_tetes, dict):
        keys = list(en_tetes.keys())
        header_labels = list(en_tetes.values())
    else:
        keys = list(en_tetes)
        header_labels = list(en_tetes)

    # Définition de la charte de styles openpyxl
    header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")  # Emerald 700
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")  # Slate 50
    data_font = Font(name="Segoe UI", size=9, color="0F172A")
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_right = Alignment(horizontal="right", vertical="center")
    data_align_center = Alignment(horizontal="center", vertical="center")

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    data_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    # 1. Écriture des en-têtes
    ws.append(header_labels)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(header_labels) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = Border(bottom=Side(border_style="medium", color="065F46"))

    # 2. Écriture des lignes de données
    for row_idx, ligne in enumerate(lignes, start=2):
        ws.row_dimensions[row_idx].height = 20
        is_alt = (row_idx % 2 == 1)
        row_values = []

        for k in keys:
            val = ligne.get(k, "")
            row_values.append(val)

        ws.append(row_values)

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = data_border
            if is_alt:
                cell.fill = row_alt_fill

            # Formatage de l'alignement selon le type de contenu
            if isinstance(val, (int, float)):
                cell.alignment = data_align_right
            elif isinstance(val, str) and (val.startswith("202") and len(val) == 10):
                cell.alignment = data_align_center
            else:
                cell.alignment = data_align_left

    # 3. Ajustement automatique de la largeur des colonnes
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    xlsx_data = output.getvalue()

    response = Response(
        xlsx_data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response


# =============================================================================
# 3. EXPORT RAPPORT DE DÉCISION PDF (Dashboard & Alertes)
# =============================================================================

def generer_rapport_dashboard_pdf(kpis, points_attention, filtres_libelles, nom_fichier="rapport_gtm_dashboard.pdf"):
    """
    Génère un rapport exécutif PDF A4 récapitulant les KPI globaux et les alertes de gestion.
    
    :param kpis: Dictionnaire des KPI calculés par stats_service
    :param points_attention: Dictionnaire des alertes générées par points_attention_service
    :param filtres_libelles: Libellé textuel du périmètre de filtre actif
    :param nom_fichier: Nom du document PDF généré
    :return: Objet Response Flask
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm
    )

    styles = getSampleStyleSheet()
    
    # Définition des styles typographiques du rapport
    style_titre = ParagraphStyle(
        'DocTitre',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A')
    )
    style_sous_titre = ParagraphStyle(
        'DocSousTitre',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#64748B')
    )
    style_section = ParagraphStyle(
        'DocSection',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#047857'),
        spaceBefore=10,
        spaceAfter=6
    )
    style_corps = ParagraphStyle(
        'DocCorps',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#334155')
    )
    style_cell_header = ParagraphStyle(
        'DocCellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )
    style_cell_bold = ParagraphStyle(
        'DocCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0F172A')
    )
    style_cell = ParagraphStyle(
        'DocCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#334155')
    )

    elements = []

    # 1. En-tête officiel du Document
    now_str = datetime.now().strftime("%d/%m/%Y à %H:%M")
    header_data = [
        [
            Paragraph("<b>GALAXY TRAINING MANAGER (GTM)</b>", ParagraphStyle('H1', fontName='Helvetica-Bold', fontSize=12, textColor=colors.HexColor('#047857'))),
            Paragraph(f"Édité le : <b>{now_str}</b>", ParagraphStyle('H2', fontName='Helvetica', fontSize=8, alignment=2, textColor=colors.HexColor('#64748B')))
        ],
        [
            Paragraph("Rapport de Synthèse & Pilotage Décisionnel", style_titre),
            Paragraph("Périmètre : " + (filtres_libelles or "Global Entreprise"), style_sous_titre)
        ]
    ]
    t_header = Table(header_data, colWidths=[12 * cm, 6 * cm])
    t_header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(t_header)
    elements.append(Spacer(1, 0.4 * cm))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#047857'), spaceBefore=2, spaceAfter=12))

    # 2. Tableau des Indicateurs Clés (KPIs)
    elements.append(Paragraph("1. Indicateurs Clés de Performance (KPIs)", style_section))

    taux_remplissage = kpis.get("taux_moyen_remplissage", 0)
    taux_str = f"{taux_remplissage:.1f}%" if isinstance(taux_remplissage, (int, float)) else "–"

    kpi_data = [
        [
            Paragraph("Sessions Actives", style_cell_header),
            Paragraph("Taux de Remplissage", style_cell_header),
            Paragraph("Participants Distincts", style_cell_header),
            Paragraph("Clients Actifs", style_cell_header),
            Paragraph("Formations", style_cell_header),
            Paragraph("Formateurs", style_cell_header)
        ],
        [
            Paragraph(str(kpis.get("sessions_actives", 0)), style_cell_bold),
            Paragraph(taux_str, style_cell_bold),
            Paragraph(str(kpis.get("participants_distincts", 0)), style_cell_bold),
            Paragraph(str(kpis.get("clients_actifs", 0)), style_cell_bold),
            Paragraph(str(kpis.get("formations_catalogue", 0)), style_cell_bold),
            Paragraph(str(kpis.get("formateurs_mobilises", 0)), style_cell_bold)
        ]
    ]
    t_kpi = Table(kpi_data, colWidths=[3 * cm] * 6)
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#047857')),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#F8FAFC')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t_kpi)
    elements.append(Spacer(1, 0.5 * cm))

    # 3. Points d'Attention & Alertes Métier
    elements.append(Paragraph("2. Points d'Attention & Alertes de Gestion", style_section))

    items_alertes = points_attention.get("items", [])
    if not items_alertes:
        elements.append(Paragraph("<i>Aucun point d'attention critique détecté sur ce périmètre. Les indicateurs sont nominaux.</i>", style_corps))
    else:
        alert_data = [
            [
                Paragraph("Niveau", style_cell_header),
                Paragraph("Sujet / Titre", style_cell_header),
                Paragraph("Détail & Recommandation", style_cell_header)
            ]
        ]
        for a in items_alertes[:12]:  # Affichage des 12 premières alertes
            niveau = a.get("niveau", "info")
            if niveau == "danger":
                badge_text = "<font color='#991B1B'><b>CRITIQUE</b></font>"
            elif niveau == "warning":
                badge_text = "<font color='#92400E'><b>ATTENTION</b></font>"
            else:
                badge_text = "<font color='#075985'><b>INFO</b></font>"

            alert_data.append([
                Paragraph(badge_text, style_cell),
                Paragraph(f"<b>{a.get('titre', '')}</b>", style_cell_bold),
                Paragraph(a.get("message", ""), style_cell)
            ])

        t_alertes = Table(alert_data, colWidths=[2.5 * cm, 5.5 * cm, 10 * cm])
        t_alertes.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        elements.append(t_alertes)

    elements.append(Spacer(1, 0.6 * cm))

    # 4. Pied de page
    elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#CBD5E1'), spaceBefore=8, spaceAfter=8))
    elements.append(Paragraph(
        "Ce document est une synthèse décisionnelle automatisée issue du progiciel Galaxy Training Manager (GTM). "
        "Il intègre les règles de gestion d'activités clients (6 mois) et de seuils de remplissage.",
        ParagraphStyle('FooterNote', fontName='Helvetica-Oblique', fontSize=7.5, leading=10, textColor=colors.HexColor('#94A3B8'))
    ))

    doc.build(elements)
    pdf_data = buffer.getvalue()

    response = Response(pdf_data, mimetype="application/pdf")
    response.headers["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response


# =============================================================================
# 4. EXPORT FICHE SESSION / FEUILLE D'ÉMARGEMENT PDF
# =============================================================================

def generer_fiche_session_pdf(session_dict, participants_inscrits, nom_fichier=None):
    """
    Génère la feuille d'émargement officielle d'une session au format PDF A4.
    Comprend le tableau complet des participants inscrits avec colonnes de signature matin et après-midi.
    
    :param session_dict: Dictionnaire décrivant la session (dates, formateur, lieu, etc.)
    :param participants_inscrits: Liste des inscriptions confirmées
    :param nom_fichier: Nom du fichier généré
    :return: Objet Response Flask
    """
    if not nom_fichier:
        nom_fichier = f"session_{session_dict.get('id', 'details')}_emargement.pdf"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm
    )

    styles = getSampleStyleSheet()

    style_titre = ParagraphStyle(
        'DocTitreSession',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0F172A')
    )
    style_cell_header = ParagraphStyle(
        'DocCellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )
    style_cell = ParagraphStyle(
        'DocCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )
    style_section = ParagraphStyle(
        'DocSection',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#047857'),
        spaceBefore=8,
        spaceAfter=4
    )

    elements = []

    # En-tête de la session
    elements.append(Paragraph("<b>GALAXY TRAINING MANAGER</b> • Feuille d'Émargement & Fiche Session", ParagraphStyle('PHeader', fontName='Helvetica-Bold', fontSize=10, textColor=colors.HexColor('#047857'))))
    elements.append(Spacer(1, 0.2 * cm))
    elements.append(Paragraph(session_dict.get('formation', {}).get('titre', 'Formation Professionnelle'), style_titre))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#047857'), spaceBefore=4, spaceAfter=8))

    # Détails logistiques de la session
    info_data = [
        [
            Paragraph(f"<b>Session N° :</b> {session_dict.get('id', '')}", style_cell),
            Paragraph(f"<b>Dates :</b> Du {session_dict.get('date_debut', '')} au {session_dict.get('date_fin', '')}", style_cell),
        ],
        [
            Paragraph(f"<b>Formateur :</b> {session_dict.get('formateur', {}).get('nom', 'Non assigné')}", style_cell),
            Paragraph(f"<b>Lieu :</b> {session_dict.get('lieu', '')} ({session_dict.get('type', '').upper()})", style_cell),
        ],
        [
            Paragraph(f"<b>Statut :</b> {session_dict.get('statut', '').upper()}", style_cell),
            Paragraph(f"<b>Capacité / Remplissage :</b> {session_dict.get('nb_inscrits_confirmes', 0)} / {session_dict.get('capacite_max', 0)} places ({(session_dict.get('taux_remplissage', 0) * 100):.0f}%)", style_cell),
        ]
    ]
    t_info = Table(info_data, colWidths=[9 * cm, 9 * cm])
    t_info.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t_info)
    elements.append(Spacer(1, 0.4 * cm))

    # Tableau des Participants avec zones d'émargement
    elements.append(Paragraph(f"Liste des Participants Inscrits ({len(participants_inscrits)})", style_section))

    table_part = [
        [
            Paragraph("N°", style_cell_header),
            Paragraph("Nom du Participant", style_cell_header),
            Paragraph("Entreprise Cliente", style_cell_header),
            Paragraph("Statut", style_cell_header),
            Paragraph("Émargement Matin", style_cell_header),
            Paragraph("Émargement Après-midi", style_cell_header),
        ]
    ]

    for idx, p in enumerate(participants_inscrits, start=1):
        statut_label = p.get("statut", "confirmee").replace("_", " ").title()
        table_part.append([
            Paragraph(str(idx), style_cell),
            Paragraph(f"<b>{p.get('participant', {}).get('nom', '')}</b><br/>{p.get('participant', {}).get('email', '')}", style_cell),
            Paragraph(p.get('participant', {}).get('client', {}).get('nom_entreprise', 'Indépendant'), style_cell),
            Paragraph(statut_label, style_cell),
            Paragraph("", style_cell),  # Zone de signature matinale
            Paragraph("", style_cell)   # Zone de signature de l'après-midi
        ])

    t_participants = Table(table_part, colWidths=[1 * cm, 5 * cm, 4.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm])
    t_participants.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#047857')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t_participants)

    doc.build(elements)
    pdf_data = buffer.getvalue()

    response = Response(pdf_data, mimetype="application/pdf")
    response.headers["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response

